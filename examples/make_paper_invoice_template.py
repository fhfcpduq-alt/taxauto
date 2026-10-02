"""종이세금계산서 입력양식(사무실 자체) 빈 엑셀 생성.

열 정의는 config/columns.yaml 의 kinds.paper_invoice.template 을 따른다(양식과 파서가 항상 일치).
    python examples/make_paper_invoice_template.py [저장경로]
기본 저장: ./종이세금계산서_입력양식.xlsx  → 채워서 inbox/{period}/{client_id}/ 에 넣으면 된다.
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import yaml
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]


def template_spec(config_dir: Path = ROOT / "config") -> dict:
    d = yaml.safe_load((config_dir / "columns.yaml").read_text(encoding="utf-8"))
    return d["kinds"]["paper_invoice"]["template"]


def build_template(path: Path, rows: list[list] | None = None, config_dir: Path = ROOT / "config") -> Path:
    """빈 양식(또는 rows 를 채운 양식) 저장. 테스트 픽스처도 이 함수를 쓴다."""
    spec = template_spec(config_dir)
    cols = spec["columns"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "입력"
    ws.append([spec["title"]])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([spec.get("guide", "")])
    ws.append(cols)
    for c in ws[3]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
        c.alignment = Alignment(horizontal="center")
    for i, name in enumerate(cols, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(10, len(name) * 2 + 4)
    letter = {name: openpyxl.utils.get_column_letter(i) for i, name in enumerate(cols, 1)}
    for name, choices in (("구분", "매출,매입"), ("증빙종류", "세금계산서,계산서"), ("영세율", "Y,N"), ("수정", "Y,N")):
        if name in letter:
            dv = DataValidation(type="list", formula1=f'"{choices}"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{letter[name]}4:{letter[name]}500")
    for r in rows or []:
        ws.append(r)
    ws.freeze_panes = "A4"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("종이세금계산서_입력양식.xlsx")
    print(build_template(out))
