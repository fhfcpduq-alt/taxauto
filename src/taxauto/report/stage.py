"""report 단계: 거래처별 검토서·안내문 초안·수정목록 + 회차 단위 사무실 리포트.

  data/{period}/{client_id}/report/
    review.html        인쇄용 한 장 검토서 (결론 → 신고서 라인 → 검토항목 → 판정근거)
    kakao.txt          거래처 안내 초안 (차단 미해결이면 '초안 아님—발송 금지')
    review_items.xlsx  검토항목 + 해당 거래 상세 (위하고 반영·수정 목록)
  data/{period}/_dashboard.html, _briefing.md   → office.build_office_reports

숫자는 전부 엔진 결과 파일(return.json, review.json, transactions.json) 값만 쓴다.
"""

from __future__ import annotations

from collections import Counter
from datetime import date, datetime
from pathlib import Path

from ..context import RunContext, StageResult
from ..models import (
    Client,
    Direction,
    Filing,
    Line,
    ReviewItem,
    ReviewStatus,
    Severity,
    Transaction,
    VatReturn,
)
from ..redact import redact
from ..workspace import Workspace
from .fmt import (
    ALWAYS_LINES,
    BASE_CSS,
    LINE_LABELS,
    TOTAL_LINES,
    dday,
    esc,
    kdate,
    line_no,
    num,
    tax_verdict,
    won,
)
from .office import build_office_reports  # noqa: F401  (pipeline 이 report.stage 에서 import)

SEV_ORDER = {Severity.BLOCKER: 0, Severity.WARN: 1, Severity.INFO: 2}
SEV_CHIP = {Severity.BLOCKER: "red", Severity.WARN: "amber", Severity.INFO: "gray"}
DECIDER_LABEL = {"rule": "규칙", "memory": "거래처 메모리", "llm": "AI", "human": "사람", "default": "기본값"}


def run(ctx: RunContext) -> StageResult:
    ws = ctx.workspace
    ret = ws.load_return()
    if ret is None:
        return StageResult(ok=False, message="return.json 없음 — compute 단계 결과 확인")
    paths = write_client_reports(
        ws, ctx.client, ctx.filing, ret, ws.load_review(), ws.load_transactions(), ctx.today,
        office_name=str((ctx.policy.get("office") or {}).get("name") or ""),
    )
    items = ws.load_review()
    open_blk = sum(1 for i in items if i.status == ReviewStatus.OPEN and i.severity == Severity.BLOCKER)
    return StageResult(ok=True, message=f"리포트 {len(paths)}개 작성", counts={"files": len(paths), "blocker_open": open_blk})


def sorted_items(items: list[ReviewItem]) -> list[ReviewItem]:
    return sorted(items, key=lambda i: (i.status != ReviewStatus.OPEN, SEV_ORDER.get(i.severity, 9), -abs(i.tax_impact)))


def open_counts(items: list[ReviewItem]) -> dict[Severity, int]:
    c = Counter(i.severity for i in items if i.status == ReviewStatus.OPEN)
    return {s: c.get(s, 0) for s in Severity}


def write_client_reports(
    ws: Workspace,
    client: Client,
    filing: Filing,
    ret: VatReturn | None,
    items: list[ReviewItem],
    txns: list[Transaction],
    today: date,
    office_name: str = "",
) -> list[Path]:
    out = ws.report_dir
    out.mkdir(parents=True, exist_ok=True)
    state = ws.load_state() or {}
    wehago = ws.load_json("wehago_return.json")
    files = []
    p = out / "review.html"
    p.write_text(render_review_html(client, filing, ret, items, txns, today, state=state, wehago_return=wehago), encoding="utf-8")
    files.append(p)
    p = out / "kakao.txt"
    p.write_text(render_kakao(client, filing, ret, items, office_name=office_name), encoding="utf-8")
    files.append(p)
    p = out / "review_items.xlsx"
    write_review_xlsx(p, client, filing, items, txns)
    files.append(p)
    return files


# ---------------------------------------------------------------------------
# review.html
# ---------------------------------------------------------------------------


def _final_tax(ret: VatReturn | None) -> int | None:
    if ret is None or "FINAL" not in ret.lines:
        return None
    return int(ret.lines["FINAL"].tax)


def _status_chip(counts: dict[Severity, int], ret: VatReturn | None) -> str:
    if ret is None:
        return '<span class="chip red">신고서 미산출</span>'
    if counts[Severity.BLOCKER]:
        return f'<span class="chip red">신고 불가 · 차단 {counts[Severity.BLOCKER]}건 미해결</span>'
    if counts[Severity.WARN]:
        return f'<span class="chip amber">확인 후 신고 · 경고 {counts[Severity.WARN]}건</span>'
    return '<span class="chip green">검토 완료 · 신고 가능</span>'


def _line_rows(ret: VatReturn) -> str:
    rows = []
    for ln in Line:
        lv = ret.lines.get(ln.name)
        if lv is None and ln.name not in ALWAYS_LINES:
            continue
        amount, tax = (lv.amount, lv.tax) if lv else (0, 0)
        if not amount and not tax and ln.name not in ALWAYS_LINES:
            continue
        cls = ' class="total"' if ln.name in TOTAL_LINES else ""
        cnt = f'<span class="muted"> · {lv.count:,}건</span>' if lv and lv.count and ln.name not in TOTAL_LINES else ""
        rows.append(
            f"<tr{cls}><td class='mono'>({esc(ln.value)})</td><td>{esc(LINE_LABELS.get(ln.name, ln.name))}{cnt}</td>"
            f"<td class='num'>{num(amount) if amount else ''}</td><td class='num'>{num(tax)}</td></tr>"
        )
    return "\n".join(rows)


def _item_html(it: ReviewItem, tx_by_id: dict[str, Transaction]) -> str:
    detail = redact(it.detail)
    if len(detail) > 400:
        detail = detail[:400] + "…"
    txs = [tx_by_id[t] for t in it.tx_ids if t in tx_by_id]
    tx_line = ""
    if txs:
        sample = ", ".join(f"{t.tx_date.month}/{t.tx_date.day} {redact(t.counterparty_name) or '-'} {t.supply_amount:,}" for t in txs[:3])
        more = f" 외 {len(txs) - 3}건" if len(txs) > 3 else ""
        tx_line = f"<div class='muted'>거래 {len(txs)}건: {esc(sample)}{more}</div>"
    elif it.tx_ids:
        tx_line = f"<div class='muted'>거래 {len(it.tx_ids)}건</div>"
    action = f"<div><b>조치</b> {esc(it.suggested_action)}</div>" if it.suggested_action else ""
    impact = won(it.tax_impact, sign=True) if it.tax_impact else ""
    return (
        f"<tr><td><span class='chip {SEV_CHIP.get(it.severity, 'gray')}'>{esc(it.severity.value)}</span></td>"
        f"<td><b>{esc(redact(it.title))}</b>"
        f"{'<div>' + esc(detail) + '</div>' if detail else ''}{action}{tx_line}"
        f"<div class='mono muted'>{esc(it.code)} · {esc(it.id)}</div></td>"
        f"<td class='num'>{impact}</td></tr>"
    )


def _basis_html(ret: VatReturn | None, txns: list[Transaction], state: dict, wehago_return: dict | None) -> str:
    blocks = []
    purchases = [t for t in txns if t.direction == Direction.PURCHASE and t.classification]
    if purchases:
        by = Counter(t.classification.decided_by.value for t in purchases)
        need = sum(1 for t in purchases if t.classification.needs_review)
        parts = [f"{DECIDER_LABEL.get(k, k)} {v:,}" for k, v in by.most_common()]
        blocks.append(
            f"<div><b>매입 분류 {len(purchases):,}건</b><div class='muted'>{esc(' · '.join(parts))}"
            f"{f' · 검토필요 {need:,}' if need else ''}</div></div>"
        )
    if ret is not None:
        if ret.non_deductible_breakdown:
            li = "".join(
                f"<li>{esc(k)} <span class='num'>{num(v.tax)}</span></li>"
                for k, v in sorted(ret.non_deductible_breakdown.items(), key=lambda kv: -kv[1].tax)
            )
            blocks.append(f"<div><b>불공제(16) 사유별 세액</b><ul>{li}</ul></div>")
        if ret.other_deductible_breakdown:
            li = "".join(f"<li>{esc(k)} <span class='num'>{num(v.tax)}</span></li>" for k, v in ret.other_deductible_breakdown.items())
            blocks.append(f"<div><b>그 밖의 공제(14) 내역</b><ul>{li}</ul></div>")
        if ret.card_receipt_summary:
            li = "".join(
                f"<li>{esc(k)} {v.count:,}건 <span class='num'>{num(v.tax)}</span></li>" for k, v in ret.card_receipt_summary.items()
            )
            blocks.append(f"<div><b>카드·현금영수증 수령명세</b><ul>{li}</ul></div>")
    if wehago_return:
        try:
            wv = VatReturn.from_dict(wehago_return)
            w_final = wv.lines.get("FINAL")
            e_final = _final_tax(ret)
            if w_final is not None and e_final is not None:
                diff = e_final - w_final.tax
                verdict = "일치" if diff == 0 else f"차이 {won(diff, sign=True)}"
                blocks.append(f"<div><b>위하고 신고서 대사</b><div>위하고 {won(w_final.tax)} · 엔진 {won(e_final)} · {esc(verdict)}</div></div>")
        except Exception:
            pass
    unv = state.get("unverified_law_params") or []
    if unv:
        blocks.append(f"<div><b>미검증 세법 파라미터 사용 {len(unv)}개</b><div class='mono muted'>{esc(', '.join(unv[:8]))}</div></div>")
    notes = ""
    if ret is not None and ret.notes:
        notes = "<ul class='notes'>" + "".join(f"<li>{esc(redact(n))}</li>" for n in ret.notes[:12]) + "</ul>"
    grid = "".join(f"<div class='card'>{b}</div>" for b in blocks)
    return f"<div class='grid'>{grid}</div>{notes}"


REVIEW_CSS = """
.head{display:flex;justify-content:space-between;align-items:flex-end;gap:12px;border-bottom:1px solid var(--line);padding-bottom:10px}
.head .kicker{font-size:11px;font-weight:600;color:var(--accent);letter-spacing:.08em}
.head .meta{text-align:right;font-size:12px;color:var(--muted)}
.verdict{display:flex;flex-wrap:wrap;gap:16px 32px;align-items:center;margin-top:14px;padding:16px 18px;border:1px solid var(--line);border-left:5px solid var(--accent);border-radius:6px;background:var(--soft)}
.verdict .label{font-size:12px;color:var(--muted)}
.verdict .big{font-size:26px;font-weight:700;font-variant-numeric:tabular-nums}
.verdict .due{font-size:15px;font-weight:600}
.verdict .chips{margin-left:auto}
.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(240px,1fr));gap:10px}
.card{border:1px solid var(--line);border-radius:6px;padding:8px 10px}
.card ul{margin:4px 0 0;padding-left:18px}.card li .num{float:right}
ul.notes{margin:8px 0 0;padding-left:18px;color:var(--muted);font-size:12px}
.resolved td{color:var(--muted);font-size:12px}
.foot{margin-top:24px;font-size:11px;color:var(--muted);border-top:1px solid var(--line);padding-top:6px}
@media (max-width:640px){.head{flex-direction:column;align-items:flex-start}.head .meta{text-align:left}.verdict .chips{margin-left:0}}
"""


def render_review_html(
    client: Client,
    filing: Filing,
    ret: VatReturn | None,
    items: list[ReviewItem],
    txns: list[Transaction],
    today: date,
    state: dict | None = None,
    wehago_return: dict | None = None,
) -> str:
    state = state or {}
    counts = open_counts(items)
    ft = _final_tax(ret)
    label, amount = tax_verdict(ft)
    d = (filing.due_date - today).days
    tx_by_id = {t.id: t for t in txns}
    s_items = sorted_items(items)
    open_items = [i for i in s_items if i.status == ReviewStatus.OPEN]
    done_items = [i for i in s_items if i.status != ReviewStatus.OPEN]

    if open_items:
        rows = "\n".join(_item_html(i, tx_by_id) for i in open_items)
        review_html = f"<table><thead><tr><th style='width:56px'>구분</th><th>내용</th><th class='num' style='width:110px'>세액영향</th></tr></thead><tbody>{rows}</tbody></table>"
    else:
        review_html = "<p class='muted'>미해결 검토항목 없음.</p>"
    if done_items:
        li = "".join(
            f"<tr class='resolved'><td>{esc(i.status.value)}</td><td>{esc(redact(i.title))}"
            f"{' — ' + esc(redact(i.resolution)) if i.resolution else ''}{' (' + esc(i.resolved_by) + ')' if i.resolved_by else ''}</td></tr>"
            for i in done_items
        )
        review_html += f"<table style='margin-top:8px'><tbody>{li}</tbody></table>"

    lines_html = (
        "<table><thead><tr><th style='width:56px'>칸</th><th>항목</th><th class='num'>금액</th><th class='num'>세액</th></tr></thead>"
        f"<tbody>{_line_rows(ret)}</tbody></table>" if ret is not None else "<p class='muted'>신고서 미산출 — 실행 상태 확인 필요.</p>"
    )
    prelim = ""
    if filing.preliminary_notice_tax:
        prelim = f"<div class='label'>예정고지 {won(filing.preliminary_notice_tax)} 차감 반영</div>"
    cov = f"{filing.coverage_start.isoformat()} ~ {filing.coverage_end.isoformat()}"
    generated = datetime.now().strftime("%Y-%m-%d %H:%M")

    return f"""<!doctype html>
<html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{esc(client.name)} {esc(filing.period.label)} 부가세 검토서</title>
<style>{BASE_CSS}{REVIEW_CSS}</style></head>
<body><div class="page">
<div class="head">
  <div><div class="kicker">부가가치세 검토서 · {esc(filing.period.label)}</div><h1>{esc(client.name)} <span class="muted mono">{esc(client.id)}</span></h1></div>
  <div class="meta">집계 {esc(cov)}{' · 예정신고분 제외' if filing.period.kind == 'F' and filing.filed_preliminary else ''}<br>작성 {esc(generated)}</div>
</div>
<div class="verdict">
  <div><div class="label">{esc(label)}</div><div class="big">{won(amount) if amount is not None else '-'}</div>{prelim}</div>
  <div><div class="label">신고·납부기한</div><div class="due">{esc(kdate(filing.due_date, with_year=True))} · {esc(dday(d))}</div></div>
  <div class="chips">{_status_chip(counts, ret)}</div>
</div>
<h2>신고서 주요 라인</h2>
{lines_html}
<h2>검토항목 (차단 {counts[Severity.BLOCKER]} · 경고 {counts[Severity.WARN]} · 참고 {counts[Severity.INFO]})</h2>
{review_html}
<h2>판정 근거 요약</h2>
{_basis_html(ret, txns, state, wehago_return)}
<div class="foot">taxauto 독립 재계산 결과. 전자신고 제출은 담당 세무사 승인 후 사람이 진행.</div>
</div></body></html>
"""


# ---------------------------------------------------------------------------
# kakao.txt
# ---------------------------------------------------------------------------


def render_kakao(client: Client, filing: Filing, ret: VatReturn | None, items: list[ReviewItem], office_name: str = "") -> str:
    counts = open_counts(items)
    ft = _final_tax(ret)
    head: list[str] = []
    if ret is None or ft is None:
        head.append("※ 초안 아님 — 발송 금지 (신고서 미산출)")
    elif counts[Severity.BLOCKER]:
        head.append(f"※ 초안 아님 — 발송 금지 (차단 {counts[Severity.BLOCKER]}건 미해결)")
        head.append("※ 아래 금액은 확정 전입니다. 차단 항목 정리 후 다시 만드세요.")
    elif counts[Severity.WARN]:
        head.append(f"※ [내부] 경고 {counts[Severity.WARN]}건 미확인 — 확인 후 발송. 이 줄부터 구분선까지 지우고 보내세요.")
    if head:
        head.append("──────────")

    who = f"{client.contact_name}님" if client.contact_name else ("담당자님" if client.is_corporation else "대표님")
    kind = "예정" if filing.period.kind == "P" else "확정"
    period_txt = f"{filing.period.year}년 {filing.period.half}기 {kind}"
    due = kdate(filing.due_date)

    body = [f"{who}, 안녕하세요{', ' + office_name + '입니다' if office_name else ''}.",
            f"{client.name} {period_txt} 부가세 정리가 끝나서 금액 알려드립니다.", ""]
    if ft is None:
        body += ["납부할 세액: (산출 전)", f"신고·납부기한: {due}까지"]
    elif ft > 0:
        body += [f"납부할 세액: {ft:,}원", f"납부기한: {due}까지"]
        if filing.preliminary_notice_tax:
            body.append(f"(예정고지로 내신 {filing.preliminary_notice_tax:,}원은 빼고 계산한 금액입니다)")
        body += ["", "신고는 기한 전에 저희가 마치고, 납부서는 신고 후 바로 보내드릴게요."]
    elif ft < 0:
        body += [f"환급 예정 세액: {-ft:,}원", f"신고기한: {due}"]
        body += ["", "이번에는 낼 세금 없이 환급이 나옵니다. 환급 일정은 신고 후 따로 알려드릴게요."]
    else:
        body += ["이번에 납부할 세액은 없습니다.", f"신고기한: {due}", "", "신고는 기한 전에 저희가 마무리하겠습니다."]
    body += ["빠진 매출·매입 자료나 확인하실 내용 있으면 편하게 말씀 주세요."]
    return "\n".join(head + body) + "\n"


# ---------------------------------------------------------------------------
# review_items.xlsx
# ---------------------------------------------------------------------------

ITEM_COLS = [("검토ID", 14), ("구분", 6), ("상태", 9), ("코드", 22), ("내용", 40), ("상세", 50), ("세액영향", 12),
             ("권장조치", 36), ("거래수", 7), ("처리메모", 30)]
TX_COLS = [("검토ID", 14), ("구분", 6), ("거래ID", 18), ("거래일자", 11), ("출처", 18), ("증빙", 10), ("거래처명", 22),
           ("거래처사업자번호", 14), ("품목/업종", 22), ("공급가액", 13), ("세액", 11), ("합계", 13), ("현재분류", 11),
           ("불공제·제외사유", 22), ("판정", 8), ("확신도", 7), ("카드", 20), ("원본파일", 24), ("행", 6),
           ("위하고 수정내용", 30), ("완료", 6)]


def _tx_row(item_id: str, sev: str, t: Transaction) -> list:
    c = t.classification
    reason = ""
    if c:
        reason = (c.non_deductible_reason.value if c.non_deductible_reason else "") or (c.exclusion_reason.value if c.exclusion_reason else "")
    return [
        item_id, sev, t.id, t.tx_date.isoformat(), t.source.value, t.doc_type.value, redact(t.counterparty_name),
        t.counterparty_biz_no, redact(t.item or t.merchant_category), int(t.supply_amount), int(t.vat), int(t.total),
        c.category.value if c else "", reason, c.decided_by.value if c else "", round(c.confidence, 2) if c else None,
        t.card_no_masked, Path(t.source_file).name if t.source_file else "", t.row_no or None, "", "",
    ]


def write_review_xlsx(path: Path, client: Client, filing: Filing, items: list[ReviewItem], txns: list[Transaction]) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F4E79")
    head_font = Font(bold=True, color="FFFFFF")
    sev_fill = {Severity.BLOCKER.value: "FDECEA", Severity.WARN.value: "FFF6DC"}

    def setup(ws, cols):
        ws.append([c for c, _ in cols])
        for i, (_, w) in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
            cell = ws.cell(row=1, column=i)
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"

    ws1 = wb.active
    ws1.title = "검토항목"
    setup(ws1, ITEM_COLS)
    s_items = sorted_items(items)
    for it in s_items:
        ws1.append([it.id, it.severity.value, it.status.value, it.code, redact(it.title), redact(it.detail),
                    int(it.tax_impact), it.suggested_action, len(it.tx_ids), it.resolution])
        r = ws1.max_row
        if it.status == ReviewStatus.OPEN and it.severity.value in sev_fill:
            ws1.cell(row=r, column=2).fill = PatternFill("solid", fgColor=sev_fill[it.severity.value])
        for col in (5, 6, 8):
            ws1.cell(row=r, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        ws1.cell(row=r, column=7).number_format = "#,##0"

    ws2 = wb.create_sheet("해당거래")
    setup(ws2, TX_COLS)
    tx_by_id = {t.id: t for t in txns}
    seen: set[str] = set()
    for it in s_items:
        if it.status != ReviewStatus.OPEN:
            continue
        for tid in it.tx_ids:
            t = tx_by_id.get(tid)
            if t:
                ws2.append(_tx_row(it.id, it.severity.value, t))
                seen.add(tid)
    # 검토항목에 안 묶였지만 분류 검토가 필요한 매입
    for t in txns:
        if t.id not in seen and t.classification and t.classification.needs_review:
            ws2.append(_tx_row("(분류검토)", Severity.WARN.value, t))
    for row in ws2.iter_rows(min_row=2):
        for idx in (9, 10, 11):
            row[idx].number_format = "#,##0"

    for ws in (ws1, ws2):
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
    info = wb.create_sheet("정보")
    info.append(["거래처", f"{client.name} ({client.id})"])
    info.append(["회차", filing.period.label])
    info.append(["집계기간", f"{filing.coverage_start.isoformat()} ~ {filing.coverage_end.isoformat()}"])
    info.append(["기한", filing.due_date.isoformat()])
    info.append(["작성", datetime.now().strftime("%Y-%m-%d %H:%M")])
    info.column_dimensions["A"].width = 10
    info.column_dimensions["B"].width = 40
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
