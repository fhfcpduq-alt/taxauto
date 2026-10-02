"""엑셀/CSV 읽기 · 헤더행 자동 탐지 · 날짜/금액 파싱.

- .xlsx(openpyxl), .xls(xlrd 있으면), .csv(utf-8-sig → cp949), HTML 위장 .xls(홈택스 일부 다운로드)
- 파일 확장자보다 실제 내용(시그니처)을 우선한다. (확장자만 .xls 인 xlsx/HTML 대비)
- 읽기 실패는 예외 대신 ParseIssue 로 돌려준다.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from ..models import ParseIssue, Severity

HEADER_SCAN_ROWS = 15
TEXT_ENCODINGS = ("utf-8-sig", "cp949")


@dataclass
class Sheet:
    name: str
    rows: list[list[Any]]


# ---------------------------------------------------------------------------
# 파일 읽기
# ---------------------------------------------------------------------------


def read_workbook(path: Path) -> tuple[list[Sheet], list[ParseIssue]]:
    """파일 → 시트 목록. 실패/미지원은 ParseIssue."""
    path = Path(path)
    name = path.name
    try:
        data = path.read_bytes()
    except OSError as e:
        return [], [ParseIssue(name, 0, f"파일 읽기 실패: {e}", Severity.WARN)]
    suffix = path.suffix.lower()
    head = data[:8]
    try:
        if head.startswith(b"PK"):
            return _read_xlsx(data), []
        if head.startswith(b"\xd0\xcf\x11\xe0"):  # OLE2 = 구형 .xls
            return _read_xls(data, name)
        if _looks_html(data):
            return _read_html(data), []
        if suffix in (".csv", ".txt", ".tsv"):
            return [_read_csv(data, name)], []
        if suffix in (".xlsx", ".xlsm", ".xls"):
            return [], [ParseIssue(name, 0, "엑셀 파일 형식이 아님(손상 또는 다른 형식) - 원본 확인 필요", Severity.WARN)]
    except _XlsUnavailable:
        return [], [ParseIssue(name, 0, "xls 변환 필요: xlrd 미설치 - 엑셀에서 .xlsx 로 다시 저장하거나 xlrd 설치", Severity.WARN)]
    except Exception as e:  # 파일 손상 등 - 전체 실행은 계속
        return [], [ParseIssue(name, 0, f"파일 해석 실패({type(e).__name__}): {e}", Severity.WARN)]
    return [], [ParseIssue(name, 0, f"지원하지 않는 파일 형식({suffix or '확장자 없음'}) - 수동 확인", Severity.WARN)]


class _XlsUnavailable(Exception):
    pass


def _trim(row: list[Any]) -> list[Any]:
    row = list(row)
    while row and (row[-1] is None or (isinstance(row[-1], str) and not row[-1].strip())):
        row.pop()
    return row


def _read_xlsx(data: bytes) -> list[Sheet]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        out = []
        for ws in wb.worksheets:
            if hasattr(ws, "reset_dimensions"):
                ws.reset_dimensions()  # 차원정보가 잘못 기록된 파일(타 프로그램 생성) 대비
            out.append(Sheet(ws.title, [_trim(r) for r in ws.iter_rows(values_only=True)]))
        return out
    finally:
        wb.close()


def _read_xls(data: bytes, name: str) -> tuple[list[Sheet], list[ParseIssue]]:
    try:
        import xlrd  # type: ignore
    except ImportError as e:
        raise _XlsUnavailable() from e
    book = xlrd.open_workbook(file_contents=data)
    out = []
    for sh in book.sheets():
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    v: Any = xlrd.xldate_as_datetime(cell.value, book.datemode)
                elif cell.ctype == xlrd.XL_CELL_NUMBER:
                    v = int(cell.value) if float(cell.value).is_integer() else cell.value
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    v = None
                else:
                    v = cell.value
                row.append(v)
            rows.append(_trim(row))
        out.append(Sheet(sh.name, rows))
    return out, []


def _decode(data: bytes) -> str:
    for enc in TEXT_ENCODINGS:
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("cp949", errors="replace")


def _read_csv(data: bytes, name: str) -> Sheet:
    text = _decode(data)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
        delim = dialect.delimiter
    except csv.Error:
        delim = "\t" if sample.count("\t") > sample.count(",") else ","
    rows = [_trim([c if c.strip() else None for c in r]) for r in csv.reader(io.StringIO(text), delimiter=delim)]
    return Sheet(Path(name).stem, rows)


def _looks_html(data: bytes) -> bool:
    s = data[:2048].lstrip(b"\xef\xbb\xbf \r\n\t").lower()
    return s.startswith(b"<") and (b"<table" in data[:200000].lower() or b"<html" in s)


class _TableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.tables: list[list[list[Any]]] = []
        self._row: list[Any] | None = None
        self._cell: list[str] | None = None
        self._span = 1

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.tables.append([])
        elif tag == "tr":
            self._row = []
        elif tag in ("td", "th"):
            self._cell = []
            try:
                self._span = max(1, int(dict(attrs).get("colspan") or 1))
            except ValueError:
                self._span = 1
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None and self._row is not None:
            txt = " ".join("".join(self._cell).split())
            self._row.append(txt or None)
            self._row.extend([None] * (self._span - 1))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            if not self.tables:
                self.tables.append([])
            self.tables[-1].append(_trim(self._row))
            self._row = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def _read_html(data: bytes) -> list[Sheet]:
    p = _TableParser()
    p.feed(_decode(data))
    return [Sheet(f"table{i + 1}", t) for i, t in enumerate(p.tables) if t]


# ---------------------------------------------------------------------------
# 헤더 탐지
# ---------------------------------------------------------------------------

_PAREN_RE = re.compile(r"\(.*?\)|\[.*?\]|（.*?）")
_HDR_STRIP_RE = re.compile(r"[\s\*·ㆍ:_\-]")


def norm_header(v: Any) -> str:
    """헤더 비교용 정규화: 괄호 내용·공백·기호 제거, 소문자."""
    if v is None:
        return ""
    s = _PAREN_RE.sub("", str(v))
    return _HDR_STRIP_RE.sub("", s).lower()


def norm_text(v: Any) -> str:
    """문맥 키워드 비교용: 공백·가운뎃점만 제거(괄호 유지), 소문자."""
    if v is None:
        return ""
    return re.sub(r"[\s·ㆍ]", "", str(v)).lower()


def _parse_alias(alias: str) -> tuple[str, int]:
    if "@" in alias:
        a, n = alias.rsplit("@", 1)
        if n.isdigit():
            return norm_header(a), int(n)
    return norm_header(alias), 0


@dataclass
class HeaderMatch:
    row_index: int                       # 0-based (rows 기준)
    columns: dict[str, int]              # field → 열 index
    headers: list[str]                   # 열별 표시용 헤더(그룹 결합)
    title_text: str                      # 헤더 위쪽 텍스트(제목·조회조건)
    header_text: str = ""

    @property
    def data_start(self) -> int:
        return self.row_index + 1


def _column_keys(rows: list[list[Any]], idx: int) -> tuple[list[dict], list[str]]:
    """열별 비교키 {plain, occ, comp} 와 표시용 헤더."""
    hdr = rows[idx]
    grp = rows[idx - 1] if idx > 0 else []
    width = max(len(hdr), len(grp))
    # 병합셀(그룹행) 앞값 채우기
    filled: list[str] = []
    last = ""
    for c in range(width):
        g = grp[c] if c < len(grp) else None
        if g is not None and str(g).strip():
            last = str(g).strip()
        filled.append(last)
    keys: list[dict] = []
    display: list[str] = []
    seen: dict[str, int] = {}
    for c in range(width):
        raw = hdr[c] if c < len(hdr) else None
        plain_src = raw if (raw is not None and str(raw).strip()) else (grp[c] if c < len(grp) else None)
        plain = norm_header(plain_src)
        occ = 0
        if plain:
            seen[plain] = seen.get(plain, 0) + 1
            occ = seen[plain]
        comp = norm_header(filled[c]) + plain if filled[c] and plain else ""
        keys.append({"plain": plain, "occ": occ, "comp": comp})
        disp = str(plain_src).strip() if plain_src is not None else ""
        if filled[c] and raw is not None and str(raw).strip() and filled[c] != disp:
            disp = f"{filled[c]}/{disp}"
        display.append(disp or f"열{c + 1}")
    return keys, display


def map_fields(keys: list[dict], field_aliases: dict[str, list[str]]) -> dict[str, int]:
    """필드 별칭 → 열 index. 필드 순서·별칭 순서 우선, 한 열은 한 필드에만."""
    used: set[int] = set()
    out: dict[str, int] = {}
    for fld, aliases in field_aliases.items():
        for alias in aliases or []:
            a, n = _parse_alias(str(alias))
            if not a:
                continue
            hit = None
            for c, k in enumerate(keys):
                if c in used or not k["plain"]:
                    continue
                if n:
                    if k["plain"] == a and k["occ"] == n:
                        hit = c
                        break
                elif k["comp"] == a or k["plain"] == a:
                    hit = c
                    break
            if hit is not None:
                out[fld] = hit
                used.add(hit)
                break
    return out


def find_header(
    rows: list[list[Any]], field_aliases: dict[str, list[str]], scan: int = HEADER_SCAN_ROWS, min_fields: int = 2
) -> HeaderMatch | None:
    """첫 scan 행 안에서 매핑되는 필드가 가장 많은 행을 헤더로."""
    best: tuple[int, int, dict, list[str]] | None = None
    for i in range(min(scan, len(rows))):
        if not rows[i]:
            continue
        keys, disp = _column_keys(rows, i)
        cols = map_fields(keys, field_aliases)
        if len(cols) >= min_fields and (best is None or len(cols) > best[0]):
            best = (len(cols), i, cols, disp)
    if best is None:
        return None
    _, i, cols, disp = best
    title = " ".join(str(c) for r in rows[:i] for c in r if c is not None and str(c).strip())
    header = " ".join(str(c) for c in rows[i] if c is not None)
    return HeaderMatch(i, cols, disp, title, header)


def is_repeated_header(row: list[Any], known_headers: set[str]) -> bool:
    """페이지마다 반복되는 헤더행인지."""
    vals = [norm_header(v) for v in row if v is not None and str(v).strip()]
    if not vals or any(not isinstance(v, str) for v in row if v is not None):
        return False
    return sum(1 for v in vals if v in known_headers) >= 2


# ---------------------------------------------------------------------------
# 값 파싱
# ---------------------------------------------------------------------------

_YMD_RE = re.compile(r"^(\d{4})\s*[-./년]\s*(\d{1,2})\s*[-./월]\s*(\d{1,2})")
_MD_RE = re.compile(r"^(\d{1,2})\s*[-./월]\s*(\d{1,2})\s*일?$")
_YM_RE = re.compile(r"^(\d{4})\s*[-./년]\s*(\d{1,2})\s*월?\.?$")
_EXCEL_EPOCH = date(1899, 12, 30)


def _empty(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() in ("", "-", "−"))


def parse_date(v: Any, default_year: int | None = None) -> date | None:
    """2026-07-03 / 20260703 / 2026.07.03 / 2026/07/03 / 2026년 7월 3일 / 엑셀 serial / datetime.
    빈값 → None, 해석 불가 → ValueError."""
    if _empty(v):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, bool):
        raise ValueError(f"날짜 아님: {v!r}")
    if isinstance(v, (int, float, Decimal)):
        if float(v).is_integer() and 19000101 <= int(v) <= 21001231:
            return _ymd(str(int(v)))
        if 20000 <= float(v) < 80000:  # 엑셀 serial (1954~2119)
            return _EXCEL_EPOCH + timedelta(days=int(v))
        raise ValueError(f"날짜 아님: {v!r}")
    s = str(v).strip()
    digits = s.replace("-", "").replace(".", "").replace("/", "").strip()
    if len(digits) == 8 and digits.isdigit():
        return _ymd(digits)
    m = _YMD_RE.match(s)
    if m:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    if s.isdigit() and len(s) == 5:
        return _EXCEL_EPOCH + timedelta(days=int(s))
    m = _MD_RE.match(s)
    if m and default_year:
        return date(default_year, int(m.group(1)), int(m.group(2)))
    raise ValueError(f"날짜 해석 불가: {s!r}")


def _ymd(s: str) -> date:
    return date(int(s[:4]), int(s[4:6]), int(s[6:8]))


def parse_month(v: Any, default_year: int | None = None) -> tuple[int, int] | None:
    """년월: 2026-07 / 2026.07 / 202607 / 2026년 7월 / datetime. 빈값 None, 불가 ValueError."""
    if _empty(v):
        return None
    if isinstance(v, (datetime, date)):
        return v.year, v.month
    s = str(v).strip()
    if isinstance(v, float) and v.is_integer():
        s = str(int(v))
    if s.isdigit() and len(s) == 6:
        y, m = int(s[:4]), int(s[4:])
    elif (mm := _YM_RE.match(s)) is not None:
        y, m = int(mm.group(1)), int(mm.group(2))
    elif s.rstrip("월").isdigit() and default_year and 1 <= int(s.rstrip("월")) <= 12:
        y, m = default_year, int(s.rstrip("월"))
    else:
        d = parse_date(v, default_year)  # 일자까지 있는 값
        return (d.year, d.month) if d else None
    if not 1 <= m <= 12:
        raise ValueError(f"년월 해석 불가: {s!r}")
    return y, m


def month_end(y: int, m: int) -> date:
    nxt = date(y + (m // 12), (m % 12) + 1, 1)
    return nxt - timedelta(days=1)


_AMT_STRIP_RE = re.compile(r"[,\s원₩\\ ]")


def parse_amount(v: Any) -> int:
    """'1,234,000' / '-' / 공백 / -1000 / '(1,000)' / '△1,000' / '1,000-' → int(원).
    소수는 반올림(엑셀 부동소수 오차 보정). 해석 불가 → ValueError."""
    if _empty(v):
        return 0
    if isinstance(v, bool):
        raise ValueError(f"금액 아님: {v!r}")
    if isinstance(v, int):
        return v
    if isinstance(v, (float, Decimal)):
        return int(Decimal(str(v)).quantize(Decimal(1), rounding=ROUND_HALF_UP))
    s = _AMT_STRIP_RE.sub("", str(v))
    if s in ("", "-", "−"):
        return 0
    neg = False
    if s.startswith("(") and s.endswith(")"):
        neg, s = True, s[1:-1]
    if s[:1] in ("△", "▲", "-", "−"):
        neg, s = not neg, s[1:]
    elif s.endswith("-"):
        neg, s = not neg, s[:-1]
    try:
        n = int(Decimal(s).quantize(Decimal(1), rounding=ROUND_HALF_UP))
    except InvalidOperation:
        raise ValueError(f"금액 해석 불가: {v!r}") from None
    return -n if neg else n


def cell_text(v: Any) -> str:
    """원본값 → 표시 문자열(raw 보존용, JSON 안전)."""
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if (v.hour, v.minute, v.second) == (0, 0, 0) else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


# ---------------------------------------------------------------------------
# 마스킹
# ---------------------------------------------------------------------------

_RRN_RE = re.compile(r"(?<!\d)(\d{6})[-\s]?([1-8])\d{6}(?!\d)")


def mask_card_no(v: Any) -> str:
    """카드번호 → 뒤 4자리만 남김: '****-****-****-3456'."""
    s = cell_text(v)
    if not s:
        return ""
    tail = re.sub(r"[^\d*]", "", s)[-4:]
    if len(tail) == 4 and tail.isdigit():
        return f"****-****-****-{tail}"
    return "****"


def mask_rrn(text: str) -> str:
    """주민등록번호 패턴 → 앞 6자리+성별자리만 남김. (raw 문자열 보존 시)"""
    if not text:
        return text
    return _RRN_RE.sub(lambda m: f"{m.group(1)}-{m.group(2)}******", text)


def is_rrn(v: Any) -> bool:
    d = "".join(ch for ch in cell_text(v) if ch.isdigit() or ch == "*")
    return len(d) == 13
