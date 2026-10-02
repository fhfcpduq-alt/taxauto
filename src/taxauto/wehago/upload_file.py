"""엔진 거래 → 위하고 매입매출전표 업로드용 xlsx.

양식(컬럼·코드)은 config/wehago/upload_template.yaml 에서 정한다(실제 위하고 양식 미확인 → 학습모드에서 확정).
계정코드·유형·적요·분개유형은 Classification 의 스타일 필드(스타일 학습 섹터가 채움)에서 가져오고,
비어 있으면 유형은 증빙·판정으로 추정하고 나머지는 빈칸(= 위하고 기본값)으로 둔다.

출력: data/{period}/{client}/wehago/upload/매입매출전표_upload.xlsx  + upload_meta.json
CLI : python -m taxauto.wehago.upload_file --client C001 --period 2026-2P [--all]
"""

from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

import yaml

from ..models import DocType, Direction, PurchaseCategory, Source, Transaction
from ..workspace import Workspace, _dump

REPO_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_TEMPLATE = REPO_ROOT / "config" / "wehago" / "upload_template.yaml"
OUT_NAME = "매입매출전표_upload.xlsx"


def load_template(path: Path | str | None = None) -> dict:
    p = Path(path) if path else DEFAULT_TEMPLATE
    t = yaml.safe_load(p.read_text(encoding="utf-8")) or {}
    if not t.get("columns"):
        raise ValueError(f"업로드 양식에 columns 가 없음: {p}")
    return t


def derive_entry_type(t: Transaction) -> str:
    """위하고 유형 이름 추정(스타일 필드가 비었을 때). 예: 과세/불공/카과/카면/현과/면세."""
    c = t.classification
    if c and c.entry_type:
        return c.entry_type
    taxed = t.vat != 0
    if t.direction == Direction.SALES:
        if t.doc_type == DocType.TAX_INVOICE:
            return "영세" if t.zero_rated else "과세"
        if t.doc_type == DocType.INVOICE:
            return "면세"
        if t.doc_type == DocType.CARD:
            return "카영" if t.zero_rated else ("카과" if taxed else "카면")
        if t.doc_type == DocType.CASH_RECEIPT:
            return "현영" if t.zero_rated else ("현과" if taxed else "현면")
        return "건별" if taxed else "면건"
    cat = c.category if c else None
    if t.doc_type == DocType.TAX_INVOICE:
        if cat == PurchaseCategory.NON_DEDUCTIBLE:
            return "불공"
        return "영세" if t.zero_rated else "과세"
    if t.doc_type == DocType.INVOICE:
        return "면세"
    deductible = taxed and cat in (PurchaseCategory.GENERAL, PurchaseCategory.FIXED_ASSET, None)
    if t.doc_type == DocType.CARD:
        return "카과" if deductible else "카면"
    if t.doc_type == DocType.CASH_RECEIPT:
        return "현과" if deductible else "현면"
    return "과세" if taxed else "면세"


def _fmt(v: Any, fmt: str | None) -> Any:
    if v is None:
        return None
    if fmt is None:
        return v
    if isinstance(v, date):
        return {"date": v, "yyyymmdd": v.strftime("%Y%m%d"), "yyyy-mm-dd": v.isoformat(),
                "mm": f"{v.month:02d}", "dd": f"{v.day:02d}"}.get(fmt, v.isoformat())
    if fmt == "biz_dash":
        d = "".join(ch for ch in str(v) if ch.isdigit())
        return f"{d[:3]}-{d[3:5]}-{d[5:]}" if len(d) == 10 else d
    if fmt == "int":
        return int(v)
    return v


def field_value(t: Transaction, field: str, tpl: dict) -> Any:
    c = t.classification
    direction = t.direction.value
    if field == "direction":
        return direction
    if field == "entry_type":
        return derive_entry_type(t)
    if field == "entry_type_code":
        name = derive_entry_type(t)
        if name.isdigit():
            return name
        return ((tpl.get("entry_type_codes") or {}).get(direction) or {}).get(name, name)
    if field in ("account_code", "account_name", "summary_text", "settlement"):
        return (getattr(c, field, "") or None) if c else None
    if field == "settlement_code":
        s = (c.settlement if c else "") or ""
        return (tpl.get("settlement_codes") or {}).get(s, s) or None
    if field == "non_deductible_reason":
        return c.non_deductible_reason.value if c and c.non_deductible_reason else None
    if field == "non_deductible_reason_code":
        if not (c and c.non_deductible_reason):
            return None
        code = (tpl.get("non_deductible_reason_codes") or {}).get(c.non_deductible_reason.value)
        return code if code is not None else c.non_deductible_reason.value
    if field == "electronic":
        return "1" if t.source in (Source.ETAX_SALES, Source.ETAX_PURCHASE, Source.EINV_SALES, Source.EINV_PURCHASE) else "0"
    if field == "summary_text_or_item":
        return ((c.summary_text if c else "") or t.item) or None
    return getattr(t, field, None)


def select_transactions(txns: list[Transaction], tpl: dict, include_all: bool = False) -> tuple[list[Transaction], list[dict]]:
    srcs = [str(x) for x in tpl.get("include_sources") or []]
    excl = {str(x) for x in tpl.get("exclude_categories") or []}
    picked, skipped = [], []
    for t in txns:
        if t.source == Source.WEHAGO_LEDGER:
            continue
        if not include_all and "ALL" not in srcs and t.source.name not in srcs:
            continue
        c = t.classification
        if c and c.category.name in excl:
            skipped.append({"tx_id": t.id, "reason": f"제외 분류 {c.category.value}"})
            continue
        picked.append(t)
    picked.sort(key=lambda x: (x.tx_date, x.direction.value, x.counterparty_name, x.id))
    return picked, skipped


def build_rows(txns: list[Transaction], tpl: dict) -> tuple[list[str], list[list[Any]]]:
    cols = tpl["columns"]
    headers = [str(c["header"]) for c in cols]
    rows = [[_fmt(field_value(t, str(c["field"]), tpl), c.get("format")) for c in cols] for t in txns]
    return headers, rows


def write_upload_xlsx(
    workspace_root: Path | str,
    *,
    template: dict | None = None,
    template_path: Path | str | None = None,
    out_path: Path | str | None = None,
    txns: list[Transaction] | None = None,
    include_all: bool = False,
) -> dict:
    """반환: {path, rows, skipped, confirmed, tx_ids}. 대상 거래가 없으면 파일을 만들지 않는다(path=None)."""
    from openpyxl import Workbook

    root = Path(workspace_root)
    tpl = template or load_template(template_path)
    all_txns = txns if txns is not None else Workspace(root).load_transactions()
    picked, skipped = select_transactions(all_txns, tpl, include_all)
    out = Path(out_path) if out_path else root / "wehago" / "upload" / OUT_NAME
    meta = {"path": None, "rows": len(picked), "skipped": skipped, "confirmed": bool(tpl.get("confirmed")),
            "tx_ids": [t.id for t in picked], "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
            "warning": None if tpl.get("confirmed") else "업로드 양식 미확정 — 사람이 위하고 양식과 대조하기 전 업로드 금지"}
    if not picked:
        return meta
    headers, rows = build_rows(picked, tpl)
    wb = Workbook()
    ws = wb.active
    ws.title = str(tpl.get("sheet_name") or "Sheet1")[:31]
    header_row = int(tpl.get("header_row") or 1)
    start_row = int(tpl.get("start_row") or header_row + 1)
    for j, h in enumerate(headers, 1):
        ws.cell(row=header_row, column=j, value=h)
    for i, r in enumerate(rows):
        for j, v in enumerate(r, 1):
            ws.cell(row=start_row + i, column=j, value=v)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    meta["path"] = str(out)
    _dump(out.parent / "upload_meta.json", meta)
    return meta


def main(argv: list[str] | None = None) -> int:
    import argparse

    ap = argparse.ArgumentParser(prog="python -m taxauto.wehago.upload_file", description="위하고 매입매출전표 업로드 xlsx 생성")
    ap.add_argument("--client", required=True)
    ap.add_argument("--period", required=True)
    ap.add_argument("--base-dir", default=".")
    ap.add_argument("--all", action="store_true", help="include_sources 무시하고 전부")
    a = ap.parse_args(argv)
    from .replay import workspace_root

    root = workspace_root(Path(a.base_dir).resolve(), a.period, a.client)
    meta = write_upload_xlsx(root, include_all=a.all)
    print(json.dumps({k: v for k, v in meta.items() if k != "tx_ids"}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
