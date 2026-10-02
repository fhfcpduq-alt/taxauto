import json
from datetime import date

from openpyxl import load_workbook
from orch_helpers import TODAY, fake_stages, make_home, sample_return, sample_review, sample_txns

from taxauto.models import Client, Filing, ReviewStatus, TaxPeriod, TaxpayerType
from taxauto.pipeline import run_period
from taxauto.report.office import build_office_reports, dashboard_order, render_briefing
from taxauto.report.stage import render_kakao, render_review_html, write_review_xlsx

P = "2026-2F"


def _client(corp=False):
    return Client(id="C9", name="테스트상회", biz_no="1234567890",
                  taxpayer_type=TaxpayerType.CORPORATION if corp else TaxpayerType.INDIVIDUAL)


def _filing(notice=0):
    return Filing(client_id="C9", period=TaxPeriod.parse(P), coverage_start=date(2026, 7, 1),
                  coverage_end=date(2026, 12, 31), due_date=date(2027, 1, 25), preliminary_notice_tax=notice)


def test_review_html_structure_and_masking():
    txns = sample_txns("C9")
    items = sample_review("C9", P, txns, blockers=1)
    html = render_review_html(_client(), _filing(), sample_return("C9", P, 295000), items, txns, TODAY)
    assert "납부할 세액" in html and "295,000원" in html
    assert "2027년 1월 25일(월)" in html and "D-115" in html
    assert "신고 불가 · 차단 1건 미해결" in html
    # 차단 → 경고 → 참고 순서
    assert html.index("차단 항목 0") < html.index("경고 항목 0") < html.index("참고 항목")
    # 주민번호 마스킹, 외부 리소스 없음
    assert "900101-1******" in html
    assert "900101-1234567" not in html
    assert "http://" not in html and "https://" not in html
    assert "접대비및이와유사한비용" in html  # 판정 근거
    assert "(27)" in html


def test_review_html_refund_and_clear():
    html = render_review_html(_client(), _filing(), sample_return("C9", P, -120000), [], [], TODAY)
    assert "환급받을 세액" in html and "120,000원" in html and "검토 완료 · 신고 가능" in html


def test_kakao_blocker_header():
    txns = sample_txns("C9")
    t = render_kakao(_client(), _filing(), sample_return("C9", P, 295000), sample_review("C9", P, txns), office_name="세무회계 민")
    assert t.startswith("※ 초안 아님 — 발송 금지 (차단 1건 미해결)")
    assert "납부할 세액: 295,000원" in t and "1월 25일(월)까지" in t


def test_kakao_clean_and_notice_tax():
    t = render_kakao(_client(), _filing(notice=500000), sample_return("C9", P, 295000), [], office_name="세무회계 민")
    assert not t.startswith("※")
    assert t.startswith("대표님, 안녕하세요, 세무회계 민입니다.")
    assert "예정고지로 내신 500,000원은 빼고" in t
    for bad in ("고객님", "드리겠사오니", "😊", "AI"):
        assert bad not in t
    corp = render_kakao(_client(corp=True), _filing(), sample_return("C9", P, -10000), [])
    assert corp.startswith("담당자님") and "환급 예정 세액: 10,000원" in corp


def test_kakao_warn_only_internal_header():
    txns = sample_txns("C9")
    items = [i for i in sample_review("C9", P, txns, blockers=0)]
    t = render_kakao(_client(), _filing(), sample_return("C9", P, 1000), items)
    assert t.startswith("※ [내부] 경고 1건 미확인")


def test_review_xlsx(tmp_path):
    txns = sample_txns("C9")
    items = sample_review("C9", P, txns)
    items[1].status = ReviewStatus.ACCEPTED
    p = write_review_xlsx(tmp_path / "r.xlsx", _client(), _filing(), items, txns)
    wb = load_workbook(p)
    ws = wb["검토항목"]
    assert ws.cell(1, 1).value == "검토ID" and ws.max_row == 1 + len(items)
    assert ws.cell(2, 2).value == "차단"
    tx = wb["해당거래"]
    rows = list(tx.iter_rows(min_row=2, values_only=True))
    assert rows[0][0] == items[0].id and rows[0][9] == 100000
    assert not any(r[0] == "(분류검토)" for r in rows)  # needs_review 거래는 이미 검토항목에 묶임
    assert len(rows) == 1  # 확인후유지 항목의 거래는 제외
    assert "위하고 수정내용" in [c.value for c in tx[1]]
    assert all("1234-****-****-5678" == r[16] or r[16] in (None, "") for r in rows)


def test_report_stage_via_pipeline_and_office(tmp_path):
    base = make_home(tmp_path)
    fns = fake_stages(final_tax={"C001": 295000, "C002": 1000}, blockers={"C001": 3, "C002": 0})
    del fns["report"]  # 실제 report 단계 사용
    s = run_period(P, base_dir=base, today=TODAY, stage_fns=fns)
    assert s["totals"]["ok"] == 2
    rdir = base / "data" / P / "C001" / "report"
    for f in ("review.html", "kakao.txt", "review_items.xlsx"):
        assert (rdir / f).exists()
    row = next(r for r in s["clients"] if r["client_id"] == "C001")
    assert row["review_html"] == "C001/report/review.html"
    c2 = next(r for r in s["clients"] if r["client_id"] == "C002")
    assert c2["ready_to_file"] is True
    dash = (base / "data" / P / "_dashboard.html").read_text(encoding="utf-8")
    assert dash.index("OO식당") < dash.index("(주)민테크")  # 차단 많은 순
    brief = (base / "data" / P / "_briefing.md").read_text(encoding="utf-8")
    assert "## 사람이 볼 것 TOP" in brief and "OO식당(C001)" in brief
    assert "차단 3건(1곳)" in brief


def test_briefing_failed_and_advice():
    s = {
        "period": P, "period_label": "2026년 2기 확정", "as_of": "2027-01-20", "due_date": "2027-01-25", "d_day": 5,
        "totals": {"clients": 3, "ok": 1, "failed": 1, "not_implemented": 1, "skipped": 0, "pending": 0,
                   "ready_to_file": 0, "with_blockers": 1, "blocker_open": 2, "warn_open": 0, "payable_total": 10, "refund_total": 0},
        "clients": [
            {"client_id": "A", "name": "가", "status": "ok", "final_tax": 10, "blocker_open": 2, "warn_open": 0, "d_day": 5,
             "top_items": [{"id": "x1", "severity": "차단", "title": "합계 불일치", "tax_impact": 3000}]},
            {"client_id": "B", "name": "나", "status": "failed", "failed_stage": "normalize", "error": "ValueError: 헤더 없음",
             "blocker_open": 0, "warn_open": 0, "d_day": 5},
            {"client_id": "C", "name": "다", "status": "not_implemented", "failed_stage": "enrich", "error": "모듈 없음",
             "blocker_open": 0, "warn_open": 0, "d_day": 5},
        ],
        "not_required": [], "last_run": {"finished_at": "2027-01-20T02:10:00+09:00", "ran": ["A", "B", "C"]},
    }
    b = render_briefing(s, top_n=3)
    assert "taxauto run --period 2026-2F --client B --from-stage normalize" in b
    assert "기한 7일 이내인데 차단이 남은 곳 1곳(가)" in b
    assert "자료 단계(collect/normalize) 실패 1곳" in b
    assert "[차단] 합계 불일치 (세액영향 +3,000원) `x1`" in b
    assert [r["client_id"] for r in dashboard_order(s["clients"])][0] == "A"


def test_build_office_reports_reads_summary(tmp_path):
    pdir = tmp_path / P
    pdir.mkdir()
    (pdir / "_summary.json").write_text(json.dumps({"period": P, "period_label": "x", "totals": {}, "clients": []}), encoding="utf-8")
    out = build_office_reports(pdir)
    assert out["dashboard"].exists() and "대상 거래처 없음" in out["dashboard"].read_text(encoding="utf-8")
